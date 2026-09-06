"""Parse Excel inventory uploads and create products (create-only + fitment update)."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.models.branch import Branch, BranchStock
from app.models.motorcycle import MotorcycleModel
from app.models.product import Product, ProductCategory
from app.schemas.product import ProductImportResponse, ProductImportRowResult
from app.services.barcodes import generate_unique_barcode
from app.services.brands import ensure_product_brand

TEMPLATE_HEADERS = [
    "barcode",
    "name",
    "brand",
    "category",
    "cost_price",
    "current_selling_price",
    "stock_qty",
    "min_stock_threshold",
    # Excel cannot multi-select in one dropdown cell — one model per column.
    "applicable_model_1",
    "applicable_model_2",
    "applicable_model_3",
    "applicable_model_4",
    "applicable_model_5",
    "applicable_model_6",
    "applicable_model_7",
    "applicable_model_8",
]

APPLICABLE_MODEL_HEADERS = [
    "applicable_model_1",
    "applicable_model_2",
    "applicable_model_3",
    "applicable_model_4",
    "applicable_model_5",
    "applicable_model_6",
    "applicable_model_7",
    "applicable_model_8",
]
APPLICABLE_MODEL_SLOTS = len(APPLICABLE_MODEL_HEADERS)
# Products sheet columns for model slots (1-based Excel col index).
APPLICABLE_MODEL_FIRST_COL = 9  # column I

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000

EXAMPLE_ROW = [
    "4800012345678",
    "Sample Brake Pad",
    "OEM",
    "Brakes",
    "150.00",
    "250.00",
    "10",
    "2",
    "Honda Click 125i",
    "Honda Wave 110",
    "",
    "",
    "",
    "",
    "",
    "",
]

# Second example: blank barcode → system generates RMS… and matches duplicates by name+brand
EXAMPLE_ROW_NO_BARCODE = [
    "",
    "Custom Labor Hose",
    "OEM",
    "Parts",
    "50.00",
    "120.00",
    "5",
    "1",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
]


def _normalize_brand_list(brand_names: list[str] | None) -> list[str]:
    brands = [b for b in (brand_names or []) if b and b.strip()]
    if "OEM" not in brands:
        # Keep example / dropdown usable even if catalog is empty.
        brands = ["OEM", *brands]
    return brands


def _normalize_category_list(category_names: list[str] | None) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for raw in category_names or []:
        name = raw.strip()
        if not name:
            continue
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        names.append(name)
    # Keep example rows valid even if catalog is empty.
    for example in ("Brakes", "Parts"):
        if example.casefold() not in seen:
            names.append(example)
            seen.add(example.casefold())
    return names


def _normalize_model_list(model_names: list[str] | None) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for raw in model_names or []:
        name = raw.strip()
        if not name:
            continue
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        names.append(name)
    # Keep examples valid even if catalog is empty.
    for example in ("Honda Click 125i", "Honda Wave 110"):
        if example.casefold() not in seen:
            names.append(example)
            seen.add(example.casefold())
    return names


def _style_products_sheet(
    ws: Worksheet,
    brand_names: list[str],
    category_names: list[str],
    model_names: list[str],
    data_row_count: int,
) -> None:
    brands_ws = ws.parent.create_sheet("Brands")
    brands_ws.append(["brand"])
    brands_ws["A1"].font = Font(bold=True)
    for name in brand_names:
        brands_ws.append([name])
    brands_ws.column_dimensions["A"].width = 20

    categories_ws = ws.parent.create_sheet("Categories")
    categories_ws.append(["category"])
    categories_ws["A1"].font = Font(bold=True)
    for name in category_names:
        categories_ws.append([name])
    categories_ws.column_dimensions["A"].width = 20

    models_ws = ws.parent.create_sheet("MotorcycleModels")
    models_ws.append(["display_name"])
    models_ws["A1"].font = Font(bold=True)
    for name in model_names:
        models_ws.append([name])
    models_ws.column_dimensions["A"].width = 28

    last_data_row = max(data_row_count, MAX_IMPORT_ROWS) + 1

    # Brand dropdown on Products!C (allow blank for unbranded items).
    last_brand_row = max(2, len(brand_names) + 1)
    dv = DataValidation(
        type="list",
        formula1=f"Brands!$A$2:$A${last_brand_row}",
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "Select a brand from the dropdown list (or leave blank)"
    dv.errorTitle = "Brand"
    dv.prompt = "Choose a seeded brand"
    dv.promptTitle = "Brand"
    ws.add_data_validation(dv)
    dv.add(f"C2:C{last_data_row}")

    # Category dropdown on Products!D (blank = no category / null).
    # showErrorMessage=False so staff can type a new category name to create on import.
    last_category_row = max(2, len(category_names) + 1)
    dv_categories = DataValidation(
        type="list",
        formula1=f"Categories!$A$2:$A${last_category_row}",
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=False,
    )
    dv_categories.prompt = (
        "Pick a category from the list, type a new name to create it, or leave blank."
    )
    dv_categories.promptTitle = "Category"
    ws.add_data_validation(dv_categories)
    dv_categories.add(f"D2:D{last_data_row}")

    # One dropdown per model slot (Excel cannot multi-select in a single cell).
    last_model_row = max(2, len(model_names) + 1)
    dv_models = DataValidation(
        type="list",
        formula1=f"MotorcycleModels!$A$2:$A${last_model_row}",
        allow_blank=True,
        showDropDown=False,
        # Allow last-slot overflow ("Model A; Model B") when a product fits >8 models.
        showErrorMessage=False,
    )
    dv_models.prompt = (
        "Pick one model per column (model 1, model 2, …). "
        "Leave unused columns blank."
    )
    dv_models.promptTitle = "Applicable model"
    ws.add_data_validation(dv_models)
    for offset in range(APPLICABLE_MODEL_SLOTS):
        col_letter = get_column_letter(APPLICABLE_MODEL_FIRST_COL + offset)
        dv_models.add(f"{col_letter}2:{col_letter}{last_data_row}")
        ws.column_dimensions[col_letter].width = 22

    # Keep barcodes as text so Excel does not strip zeros / use scientific notation.
    if data_row_count > 0:
        for row in ws.iter_rows(min_row=2, max_row=data_row_count + 1, max_col=1):
            for cell in row:
                cell.number_format = "@"

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 18


def build_import_template(
    brand_names: list[str] | None = None,
    category_names: list[str] | None = None,
    motorcycle_model_names: list[str] | None = None,
) -> bytes:
    brands = _normalize_brand_list(brand_names)
    categories = _normalize_category_list(category_names)
    models = _normalize_model_list(motorcycle_model_names)

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(EXAMPLE_ROW)
    ws.append(EXAMPLE_ROW_NO_BARCODE)
    _style_products_sheet(
        ws, brands, categories, models, data_row_count=2
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export_xlsx(
    rows: list[list[Any]],
    brand_names: list[str] | None = None,
    category_names: list[str] | None = None,
    motorcycle_model_names: list[str] | None = None,
) -> bytes:
    """Build an .xlsx using the same columns as the import template."""
    brands = _normalize_brand_list(brand_names)
    categories = _normalize_category_list(category_names)
    models = _normalize_model_list(motorcycle_model_names)
    seen_brands = {b.casefold() for b in brands}
    seen_categories = {c.casefold() for c in categories}
    seen_models = {m.casefold() for m in models}
    for row in rows:
        brand = _cell_str(row[2]) if len(row) > 2 else ""
        if brand and brand.casefold() not in seen_brands:
            brands.append(brand)
            seen_brands.add(brand.casefold())
        category = _cell_str(row[3]) if len(row) > 3 else ""
        if category and category.casefold() not in seen_categories:
            categories.append(category)
            seen_categories.add(category.casefold())
        for offset in range(APPLICABLE_MODEL_SLOTS):
            idx = 8 + offset
            if len(row) <= idx:
                break
            for part in _split_applicable_models(_cell_str(row[idx])):
                if part.casefold() not in seen_models:
                    models.append(part)
                    seen_models.add(part.casefold())

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    _style_products_sheet(
        ws, brands, categories, models, data_row_count=len(rows)
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        # Avoid scientific notation for long numeric barcodes stored as float.
        text = f"{value:.0f}" if abs(value) >= 1e11 else str(value).rstrip("0").rstrip(".")
        return text.strip()
    if isinstance(value, Decimal):
        if value == value.to_integral_value():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def _split_applicable_models(raw: str) -> list[str]:
    if not raw.strip():
        return []
    parts: list[str] = []
    seen: set[str] = set()
    for chunk in raw.replace("|", ";").split(";"):
        name = chunk.strip()
        if not name:
            continue
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        parts.append(name)
    return parts


def _dedupe_model_names(names: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for name in names:
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        out.append(name)
    return out


def _collect_applicable_model_names(col) -> list[str]:
    """Gather model names from applicable_model_1..8 and legacy applicable_models."""
    names: list[str] = []
    for header in APPLICABLE_MODEL_HEADERS:
        names.extend(_split_applicable_models(col(header)))
    # Legacy single-column format (semicolon-separated).
    names.extend(_split_applicable_models(col("applicable_models")))
    return _dedupe_model_names(names)


def models_to_export_slots(display_names: list[str]) -> list[str]:
    """Split model names into Excel dropdown columns (overflow joins into last slot)."""
    slots = [""] * APPLICABLE_MODEL_SLOTS
    if not display_names:
        return slots
    if len(display_names) <= APPLICABLE_MODEL_SLOTS:
        for i, name in enumerate(display_names):
            slots[i] = name
        return slots
    for i in range(APPLICABLE_MODEL_SLOTS - 1):
        slots[i] = display_names[i]
    slots[-1] = "; ".join(display_names[APPLICABLE_MODEL_SLOTS - 1 :])
    return slots


def _has_applicable_model_columns(header_map: dict[str, int]) -> bool:
    if "applicable_models" in header_map:
        return True
    return any(h in header_map for h in APPLICABLE_MODEL_HEADERS)


def _motorcycle_model_cache(db: Session) -> dict[str, MotorcycleModel]:
    rows = db.scalars(
        select(MotorcycleModel).where(MotorcycleModel.is_active.is_(True))
    ).all()
    return {m.display_name.casefold(): m for m in rows}


def _resolve_applicable_models(
    names: list[str],
    cache: dict[str, MotorcycleModel],
) -> list[MotorcycleModel]:
    if not names:
        return []
    missing: list[str] = []
    resolved: list[MotorcycleModel] = []
    for name in names:
        model = cache.get(name.casefold())
        if model is None:
            missing.append(name)
        else:
            resolved.append(model)
    if missing:
        raise ValueError(
            "Unknown motorcycle model(s): "
            + "; ".join(missing)
            + " — use exact names from the MotorcycleModels sheet / dropdowns"
        )
    return resolved


def _parse_decimal(raw: str, field: str) -> Decimal:
    try:
        value = Decimal(raw)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} must be a number") from exc
    if value < 0:
        raise ValueError(f"{field} must be >= 0")
    return value


def _parse_nonneg_int(raw: str, field: str, default: int = 0) -> int:
    if not raw:
        return default
    try:
        if "." in raw:
            as_dec = Decimal(raw)
            if as_dec != as_dec.to_integral_value():
                raise ValueError
            value = int(as_dec)
        else:
            value = int(raw)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} must be a whole number >= 0") from exc
    if value < 0:
        raise ValueError(f"{field} must be >= 0")
    return value


def _header_map(ws: Worksheet) -> dict[str, int]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise ValueError("Spreadsheet is empty")
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = _cell_str(cell).lower()
        if key:
            mapping[key] = idx
    missing = [
        h
        for h in ("barcode", "name", "cost_price", "current_selling_price")
        if h not in mapping
    ]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}")
    return mapping


def _get_or_create_category(
    db: Session,
    name: str,
    cache: dict[str, ProductCategory],
) -> ProductCategory:
    key = name.casefold()
    cached = cache.get(key)
    if cached is not None:
        return cached
    existing = db.scalar(
        select(ProductCategory).where(func.lower(ProductCategory.name) == key)
    )
    if existing is not None:
        cache[key] = existing
        return existing
    category = ProductCategory(name=name[:100])
    db.add(category)
    db.flush()
    cache[key] = category
    return category


def _create_product_row(
    db: Session,
    *,
    active_branch_id: UUID,
    other_branch_ids: list[UUID],
    barcode: str,
    name: str,
    brand: str | None,
    category_id: UUID | None,
    cost_price: Decimal,
    selling_price: Decimal,
    stock_qty: int,
    min_stock: int,
) -> Product:
    product = Product(
        barcode=barcode,
        name=name,
        brand=brand,
        cost_price=cost_price,
        current_selling_price=selling_price,
        stock_qty=stock_qty,
        min_stock_threshold=min_stock,
        category_id=category_id,
        is_active=True,
    )
    db.add(product)
    db.flush()
    db.add(
        BranchStock(
            branch_id=active_branch_id,
            product_id=product.id,
            stock_qty=stock_qty,
            min_stock_threshold=min_stock,
        )
    )
    for branch_id in other_branch_ids:
        db.add(
            BranchStock(
                branch_id=branch_id,
                product_id=product.id,
                stock_qty=0,
                min_stock_threshold=min_stock,
            )
        )
    return product


def _maybe_update_fitment(
    *,
    product: Product,
    model_names: list[str],
    has_models_column: bool,
    model_cache: dict[str, MotorcycleModel],
) -> str | None:
    """Update fitment when model columns are present and at least one is filled."""
    if not has_models_column or not model_names:
        return None
    product.applicable_motorcycle_models = _resolve_applicable_models(
        model_names, model_cache
    )
    return f"Fitment updated ({len(product.applicable_motorcycle_models)} model(s))"


def import_products_from_xlsx(
    db: Session,
    data: bytes,
    active_branch_id: UUID,
) -> ProductImportResponse:
    if len(data) > MAX_IMPORT_BYTES:
        raise ValueError("File too large (max 5 MB)")

    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises varied errors
        raise ValueError("Invalid Excel file — use the downloadable .xlsx template") from exc

    ws = wb.active
    header_map = _header_map(ws)
    has_models_column = _has_applicable_model_columns(header_map)
    model_cache = _motorcycle_model_cache(db)

    other_branch_ids = list(
        db.scalars(
            select(Branch.id).where(
                Branch.is_active.is_(True),
                Branch.id != active_branch_id,
            )
        ).all()
    )
    category_cache: dict[str, ProductCategory] = {}
    seen_barcodes: set[str] = set()
    seen_name_brand: set[tuple[str, str]] = set()
    results: list[ProductImportRowResult] = []
    created = updated = skipped = errors = 0
    data_rows = 0

    for excel_row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None or _cell_str(cell) == "" for cell in row):
            continue
        data_rows += 1
        if data_rows > MAX_IMPORT_ROWS:
            results.append(
                ProductImportRowResult(
                    row=excel_row_num,
                    status="error",
                    message=f"Row limit exceeded (max {MAX_IMPORT_ROWS})",
                )
            )
            errors += 1
            break

        def col(key: str) -> str:
            idx = header_map.get(key)
            if idx is None or idx >= len(row):
                return ""
            return _cell_str(row[idx])

        barcode = col("barcode")
        name = col("name")
        brand_raw = col("brand")
        category_name = col("category")
        cost_raw = col("cost_price")
        sell_raw = col("current_selling_price")
        stock_raw = col("stock_qty")
        min_raw = col("min_stock_threshold")
        model_names = _collect_applicable_model_names(col)
        auto_barcode = False

        try:
            if not name:
                raise ValueError("name is required")
            if len(name) > 200:
                raise ValueError("name must be at most 200 characters")
            if category_name and len(category_name) > 100:
                raise ValueError("category must be at most 100 characters")
            if not cost_raw:
                raise ValueError("cost_price is required")
            if not sell_raw:
                raise ValueError("current_selling_price is required")

            cost_price = _parse_decimal(cost_raw, "cost_price")
            selling_price = _parse_decimal(sell_raw, "current_selling_price")
            stock_qty = _parse_nonneg_int(stock_raw, "stock_qty", default=0)
            min_stock = _parse_nonneg_int(min_raw, "min_stock_threshold", default=0)
            brand = brand_raw[:100] if brand_raw else None
            if brand:
                ensure_product_brand(db, brand)
            applicable_models = _resolve_applicable_models(model_names, model_cache)

            name_key = name.casefold()
            brand_key = (brand or "").casefold()
            identity = (name_key, brand_key)

            if barcode:
                if len(barcode) > 64:
                    raise ValueError("barcode must be at most 64 characters")
                if barcode in seen_barcodes:
                    raise ValueError("duplicate barcode in this file")
                existing = db.scalar(select(Product).where(Product.barcode == barcode))
                if existing is not None:
                    fitment_msg = _maybe_update_fitment(
                        product=existing,
                        model_names=model_names,
                        has_models_column=has_models_column,
                        model_cache=model_cache,
                    )
                    if fitment_msg:
                        updated += 1
                        results.append(
                            ProductImportRowResult(
                                row=excel_row_num,
                                barcode=barcode,
                                name=name,
                                status="updated",
                                message=fitment_msg,
                            )
                        )
                    else:
                        skipped += 1
                        results.append(
                            ProductImportRowResult(
                                row=excel_row_num,
                                barcode=barcode,
                                name=name,
                                status="skipped",
                                message="Barcode already exists",
                            )
                        )
                    continue
            else:
                # No package barcode — match by name + brand before creating.
                if identity in seen_name_brand:
                    skipped += 1
                    results.append(
                        ProductImportRowResult(
                            row=excel_row_num,
                            barcode=None,
                            name=name,
                            status="skipped",
                            message="Duplicate name+brand in this file (no barcode)",
                        )
                    )
                    continue
                name_stmt = select(Product).where(
                    func.lower(Product.name) == name_key,
                    Product.deleted_at.is_(None),
                )
                if brand:
                    name_stmt = name_stmt.where(func.lower(Product.brand) == brand_key)
                else:
                    name_stmt = name_stmt.where(
                        or_(Product.brand.is_(None), Product.brand == "")
                    )
                existing_by_name = db.scalar(name_stmt.limit(1))
                if existing_by_name is not None:
                    fitment_msg = _maybe_update_fitment(
                        product=existing_by_name,
                        model_names=model_names,
                        has_models_column=has_models_column,
                        model_cache=model_cache,
                    )
                    if fitment_msg:
                        updated += 1
                        results.append(
                            ProductImportRowResult(
                                row=excel_row_num,
                                barcode=existing_by_name.barcode,
                                name=name,
                                status="updated",
                                message=fitment_msg,
                            )
                        )
                    else:
                        skipped += 1
                        results.append(
                            ProductImportRowResult(
                                row=excel_row_num,
                                barcode=existing_by_name.barcode,
                                name=name,
                                status="skipped",
                                message=(
                                    "Possible duplicate — same name/brand already exists "
                                    f"(barcode {existing_by_name.barcode})"
                                ),
                            )
                        )
                    continue
                barcode = generate_unique_barcode(db, reserved=seen_barcodes)
                auto_barcode = True

            seen_barcodes.add(barcode)
            seen_name_brand.add(identity)

            category_id = None
            if category_name:
                category = _get_or_create_category(db, category_name, category_cache)
                category_id = category.id
            product = _create_product_row(
                db,
                active_branch_id=active_branch_id,
                other_branch_ids=other_branch_ids,
                barcode=barcode,
                name=name,
                brand=brand,
                category_id=category_id,
                cost_price=cost_price,
                selling_price=selling_price,
                stock_qty=stock_qty,
                min_stock=min_stock,
            )
            product.applicable_motorcycle_models = applicable_models
            created += 1
            fitment_note = (
                f" · {len(applicable_models)} model(s)"
                if applicable_models
                else ""
            )
            results.append(
                ProductImportRowResult(
                    row=excel_row_num,
                    barcode=barcode,
                    name=name,
                    status="created",
                    message=(
                        ("Created (auto barcode)" if auto_barcode else "Created")
                        + fitment_note
                    ),
                )
            )
        except ValueError as exc:
            errors += 1
            results.append(
                ProductImportRowResult(
                    row=excel_row_num,
                    barcode=barcode or None,
                    name=name or None,
                    status="error",
                    message=str(exc),
                )
            )

    wb.close()

    if data_rows == 0:
        raise ValueError("No data rows found — add products below the header row")

    db.commit()
    return ProductImportResponse(
        created=created,
        updated=updated,
        skipped=skipped,
        errors=errors,
        rows=results,
    )
